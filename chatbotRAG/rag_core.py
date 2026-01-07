import torch
import os
import docx
import numpy as np
from langchain_text_splitters import RecursiveCharacterTextSplitter 
from langchain_core.documents import Document 
from langchain_community.vectorstores import Chroma
from sentence_transformers import SentenceTransformer

# --- Configuration for Chroma Persistence ---
CHROMA_PERSIST_DIR = "./chroma_db"
CHROMA_COLLECTION_NAME = "chatbot_rag_docs"

# Global state
embedding_model_rag = None
vector_store_rag = None

def get_docx_chunks(doc_path):
    """Reads a .docx file, splits its text, and returns LangChain Document objects."""
    try:
        doc = docx.Document(doc_path)
        full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    except Exception as e:
        print(f"Error reading docx file {doc_path}: {e}")
        return []

    if not full_text:
        return []

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=100,
    )
    texts = text_splitter.split_text(full_text)

    # Convert text chunks to LangChain Document objects for Chroma
    documents = [Document(page_content=t, metadata={"source": doc_path}) for t in texts]
    return documents

def setup_rag_system(document_files):
    """Initializes the embedding model and loads or creates the persistent ChromaDB vector store."""
    global embedding_model_rag, vector_store_rag

    # 1. Initialize Embedding Model (same as before)
    device = 'cuda' if torch.cuda.is_available() else 'cpu'
    embedding_model_rag = SentenceTransformer('all-MiniLM-L6-v2', device=device)

    # Helper to create a Chroma-compatible embedding function wrapper
    class SentenceTransformerEmbeddings:
        def embed_documents(self, texts):
            # Encode and convert to list format required by Chroma/LangChain
            return embedding_model_rag.encode(texts).tolist()
        def embed_query(self, text):
            return embedding_model_rag.encode(text).tolist()

    embedding_function_wrapper = SentenceTransformerEmbeddings()

    # 2. Check for existing persistent store (Decoupling Ingestion)
    if os.path.exists(CHROMA_PERSIST_DIR):
        print("Loading existing ChromaDB index...")
        try:
            vector_store_rag = Chroma(
                collection_name=CHROMA_COLLECTION_NAME,
                embedding_function=embedding_function_wrapper,
                persist_directory=CHROMA_PERSIST_DIR
            )
            print("ChromaDB index successfully loaded. (Startup is fast!)")
            return True
        except Exception as e:
             print(f"Failed to load ChromaDB: {e}. Rebuilding index.")
             # Fall through to rebuild

    # 3. Ingest documents (only if the index doesn't exist)
    all_documents = []
    for doc_file in document_files:
        print(f"Ingesting {doc_file}...")
        documents = get_docx_chunks(doc_file)
        all_documents.extend(documents)

    if not all_documents:
        print("No documents found for ingestion. RAG system setup failed.")
        return False

    print(f"Creating new ChromaDB index from {len(all_documents)} chunks. This may take a moment...")
    vector_store_rag = Chroma.from_documents(
        documents=all_documents,
        embedding=embedding_function_wrapper,
        collection_name=CHROMA_COLLECTION_NAME,
        persist_directory=CHROMA_PERSIST_DIR
    )
    vector_store_rag.persist() # Save the index to disk

    print("New ChromaDB index successfully created and saved.")
    return True

# RAG Search Tool Definition (Remains simple due to LangChain's standard interface)
def search_document_scope(query):
    """Searches the RAG documents for information."""
    global vector_store_rag

    if not vector_store_rag:
        return "Error: RAG system is not initialized."

    try:
        # Use LangChain's retriever interface for Chroma
        retriever = vector_store_rag.as_retriever(search_kwargs={"k": 3})
        # Retrieve the top 3 most relevant documents
        docs = retriever.invoke(query)
        # Combine the content of the retrieved documents for the LLM
        retrieved_context = "\n---\n".join([doc.page_content for doc in docs])
        return retrieved_context
    except Exception as e:
        return f"Error during RAG search: {e}"